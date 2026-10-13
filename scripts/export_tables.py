#!/usr/bin/env python3
"""Export the manuscript tables as LaTeX snippets (results/tables_tex) and editable
Word files (results/tables_docx/TableN.docx), from the analysis outputs."""

import argparse
import json
import os

import numpy as np
import pandas as pd

MODEL_LAB = {"llama3.1:8b": "Llama 3.1 8B", "qwen2.5:14b": "Qwen2.5 14B", "gemma3:12b": "Gemma 3 12B",
             "llama3.1_8b": "Llama 3.1 8B", "qwen2.5_14b": "Qwen2.5 14B", "gemma3_12b": "Gemma 3 12B"}
METRIC_LAB = {"fp_ratio": "First-person ratio$^\\dagger$", "self_name": "Self-reference by name$^\\dagger$",
              "subjectivity": "Subjectivity", "intensity": "Emotional intensity", "polarity": "Polarity (TextBlob)",
              "vader": "Polarity (VADER)", "distilbert": "Polarity (DistilBERT)", "n_words": "Utterance length (words)",
              "word_len": "Content-word length", "zipf": "Content-word Zipf frequency", "rare_rate": "Rare-word rate",
              "mattr": "MATTR", "mtld": "MTLD"}
TEXT_ROWS = [
    ("good_earth", "\\textit{The Good Earth}", "Buck (1931), ch. 26", "Realist family novel", "Heterodiegetic, third person"),
    ("pride_prejudice", "\\textit{Pride and Prejudice}", "Austen (1813), ch. 34", "Novel of manners", "Heterodiegetic, free indirect discourse"),
    ("great_expectations", "\\textit{Great Expectations}", "Dickens (1861), ch. 8", "Bildungsroman", "Autodiegetic, retrospective first person"),
    ("scandal_bohemia", "\\textit{A Scandal in Bohemia}", "Doyle (1891), part I", "Detective story", "Homodiegetic witness narrator"),
    ("earnest", "\\textit{The Importance of Being Earnest}", "Wilde (1895), act I", "Comic drama", "No narrator (dramatic mode)"),
    ("home_world", "\\textit{The Home and the World}", "Tagore (1916; tr. 1919), ch. 10", "Political novel", "Alternating first-person narrators"),
]


def fmt(x, d=3):
    if x is None or (isinstance(x, float) and not np.isfinite(x)):
        return "--"
    v = f"{x:.{d}f}"
    return "$-$" + v[1:] if v.startswith("-") else v  # typographic minus in LaTeX (plain '-' in Word)


def fmt_p(p):
    if p is None or not np.isfinite(p):
        return "--"
    return "$<$0.001" if p < 0.001 else f"{p:.3f}"


def latex_table(cols, rows, align):
    head = " & ".join(cols) + " \\\\"
    body = "\n".join(" & ".join(r) + " \\\\" for r in rows)
    return f"\\begin{{tabular}}{{{align}}}\n\\toprule\n{head}\n\\midrule\n{body}\n\\bottomrule\n\\end{{tabular}}\n"


def strip_tex(s):
    return (s.replace("\\textit{", "").replace("}", "").replace("$^\\dagger$", " †").replace("$<$", "<")
            .replace("$d_z$", "dz").replace("$\\Delta$", "Δ").replace("\\%", "%").replace("$", "").replace("\\", ""))


def docx_table(path, title, cols, rows, note=None):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(10)
    doc.add_paragraph().add_run(title).bold = True
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, cols):
        c.text = strip_tex(h)
        for r in c.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = strip_tex(v)
    if note:
        doc.add_paragraph(strip_tex(note))
    doc.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--res", default="results")
    ap.add_argument("--runs", default="runs")
    ap.add_argument("--excerpts", default=os.path.join(os.path.dirname(__file__), "..", "texts", "excerpts.json"))
    args = ap.parse_args()
    R = args.res
    tex_dir, doc_dir = os.path.join(R, "tables_tex"), os.path.join(R, "tables_docx")
    os.makedirs(tex_dir, exist_ok=True)
    os.makedirs(doc_dir, exist_ok=True)
    S = json.load(open(os.path.join(R, "summary.json"), encoding="utf-8"))
    ex = {t["id"]: t for t in json.load(open(args.excerpts, encoding="utf-8"))}
    tables = {}

    # Table 1: excerpts
    cols = ["Excerpt", "Source", "Genre", "Narrative structure", "Words", "Cast", "Quotations"]
    rows = [[t, s, g, n, str(ex[i]["n_words"]), str(len(ex[i]["cast"])), str(len(ex[i]["gold_quotes"]))] for i, t, s, g, n in TEXT_ROWS]
    tables[1] = ("Literary excerpts used in the experiments.", cols, rows, "p{3.0cm}p{2.6cm}p{2.2cm}p{3.0cm}rrr", None)

    # Table 2: parameters (static; mirrors the manuscript)
    tables[2] = ("Experimental parameters and configurations.", ["Parameter", "Value"], [
        ["Generator LLMs", "llama3.1:8b, qwen2.5:14b, gemma3:12b (Ollama 0.40.3, Q4_K_M)"],
        ["Judge LLMs", "phi4:14b, qwen3.5:9b-q4_K_M (temperature 0)"],
        ["Narrative modes", "first person, third person, neutral (no instruction)"],
        ["Maximum dialogue turns", "20 generated turns per session"],
        ["Generation temperature", "0.7 (Director extraction and rewriting: 0)"],
        ["Maximum tokens per response", "1000 (extraction and single-prompt baseline: 3000-4000)"],
        ["Context window", "8192 tokens"],
        ["Memory window k", "5 turns (ablation: 1)"],
        ["Speaker selection", "addressee-aware softmax, alpha = 2.0, beta = 1.5, gamma = 1.0, tau = 0.7 (ablations: uniform, round robin)"],
        ["Random seeds", "0-4 (main experiment); 0-2 (ablations and extension)"],
        ["JSON retry budget", "3 attempts per call; mode enforcement: 1 regeneration + Director rewrite"],
        ["Sentiment analysis", "DistilBERT (SST-2), TextBlob, VADER"],
        ["Subjectivity analysis", "TextBlob"],
        ["Pronoun and lexical analysis", "spaCy en_core_web_sm, wordfreq Zipf frequencies"]], "ll", None)

    # Table 3: extraction quality and pipeline reliability per model
    ext = pd.read_csv(os.path.join(R, "table_extraction.csv"))
    rel = pd.read_csv(os.path.join(R, "table_reliability.csv"))
    rel = rel[rel.condition == "full"]
    cols = ["Model", "Quote recall", "Attribution acc.", "Fabricated", "TP compliant (1st)", "TP regenerated", "TP rewritten", "Residual (TP)", "Residual (FP)"]
    rows = []
    for m in ["llama3.1:8b", "qwen2.5:14b", "gemma3:12b"]:
        e = ext[ext.model == m.replace(":", "_")]
        tp = rel[(rel.model == m) & (rel["mode"] == "third_person")]
        fp = rel[(rel.model == m) & (rel["mode"] == "first_person")]
        if tp.empty:
            continue
        rows.append([MODEL_LAB[m], fmt(e.quote_recall.mean(), 2), fmt(e.attribution_acc.mean(), 2), fmt(e.fabricated_rate.mean(), 2),
                     fmt(tp.compliant_first.iloc[0], 3), fmt(tp.regenerated.iloc[0], 3), fmt(tp.director_rewrite.iloc[0], 3),
                     fmt(tp.residual_violation.iloc[0], 3), fmt(fp.residual_violation.iloc[0], 3)])
    tables[3] = ("Director extraction quality (against manually attributed quotations, mean over six excerpts) and narrative-mode enforcement in the main experiment (share of generated turns).",
                 cols, rows, "lrrrrrrrr", None)

    # Table 5: main comparisons
    means = pd.DataFrame(S["mode_means"]).set_index("mode")
    comp = pd.DataFrame(S["mode_comparisons"])
    cc = pd.DataFrame(S["cell_consistency"]).set_index("metric")
    cols = ["Measure", "First", "Third", "Neutral", "Rule-based", "$\\Delta$ First$-$Third [95\\% CI]", "$d_z$", "$p_{\\mathrm{Holm}}$", "Cells $>0$"]
    rows = []
    for m in ["fp_ratio", "self_name", "subjectivity", "intensity", "polarity", "vader", "distilbert", "n_words", "word_len", "zipf", "rare_rate", "mattr", "mtld"]:
        r = comp[(comp.metric == m) & (comp.a == "first_person") & (comp.b == "third_person")]
        if r.empty:
            continue
        r = r.iloc[0]
        d = 1 if m in ("n_words", "mtld", "self_name") else 3
        g = lambda mode: fmt(means.loc[mode, f"{m}_mean"], d) if mode in means.index and f"{m}_mean" in means.columns else "--"
        rows.append([METRIC_LAB[m], g("first_person"), g("third_person"), g("neutral"), g("mechanical_third"),
                     f"{fmt(r['diff'], d)} [{fmt(r.ci_low, d)}, {fmt(r.ci_high, d)}]", fmt(r.dz, 2), fmt_p(r.p_holm),
                     f"{int(cc.loc[m, 'n_pos'])}/{int(cc.loc[m, 'n_cells'])}" if m in cc.index else "--"])
    tables[5] = ("Session-level means by narrative mode and paired First$-$Third comparisons (main experiment; 6 excerpts $\\times$ 3 models $\\times$ 5 seeds).",
                 cols, rows, "lrrrrcrrr",
                 "$^\\dagger$Manipulation check. Rule-based: first-person sessions transposed to third person by pronoun substitution and a reporting frame. Cells $>0$: excerpt--model cells (out of 18) in which the First$-$Third mean difference is positive. $p$ values: Wilcoxon signed-rank, Holm-adjusted across the measures of this table.")

    # Table 6: Good Earth characters
    utt = pd.read_csv(os.path.join(R, "utterances.csv"))
    u = utt[(utt.condition == "full") & (utt.text_id == "good_earth") & utt["mode"].isin(["first_person", "third_person"])]
    ident = pd.read_csv(os.path.join(R, "table_identity.csv")) if os.path.exists(os.path.join(R, "table_identity.csv")) else pd.DataFrame()
    cols = ["Character", "Mode", "Turns", "Content words", "Word length (SD)", "Zipf", "Subjectivity", "$S$ within", "$S$ between"]
    rows = []
    for c in ["O-lan", "Wang Lung", "the eldest son", "the maid"]:
        for mode, ml in (("first_person", "First"), ("third_person", "Third")):
            s = u[(u.speaker == c) & (u["mode"] == mode)]
            if s.empty:
                continue
            toks = " ".join(s.content_tokens.fillna("").astype(str)).split()
            L = np.array([len(w) for w in toks]) if toks else np.array([np.nan])
            iw = ident[(ident.text_id == "good_earth") & (ident.character == c)] if not ident.empty else pd.DataFrame()
            rows.append([c if mode == "first_person" else "", ml, str(len(s)), str(len(toks)), f"{np.nanmean(L):.2f} ({np.nanstd(L):.2f})",
                         fmt(s.zipf.mean(), 2), fmt(s.subjectivity.mean(), 3),
                         fmt(iw.within_score.mean(), 3) if mode == "first_person" and not iw.empty else "",
                         fmt(iw.between_score.mean(), 3) if mode == "first_person" and not iw.empty else ""])
    tables[6] = ("Character-level lexical statistics for \\textit{The Good Earth} (all models and seeds pooled).", cols, rows, "llrrrrrrr",
                 "$S$ within: composite consistency of the same character between first- and third-person sessions; $S$ between: mean consistency with the other characters in the same mode (averaged over models).")

    # Table 4: example of generated dialogue (same scene, model and seed in both modes)
    ex_rows = []
    for mode, lab in (("first_person", "First person"), ("third_person", "Third person")):
        p = os.path.join(args.runs, "sessions", "full", "qwen2.5_14b", "good_earth", f"{mode}_s0.json")
        if os.path.exists(p):
            r = json.load(open(p, encoding="utf-8"))
            for t in r["turns"][:8]:
                ex_rows.append([lab if t["turn"] == 1 else "", str(t["turn"]), t["speaker"], t["text"].replace("&", "\\&").replace("%", "\\%")])
    if ex_rows:
        tables[4] = ("First eight generated turns of the \\textit{The Good Earth} scene in the two narrative modes (Qwen2.5 14B, seed 0, same role cards and opening).",
                     ["Mode", "Turn", "Speaker", "Generated turn"], ex_rows, "llp{1.8cm}p{9.0cm}", None)

    # Table 7: ablations (each condition vs. the matched full-framework sessions)
    am = pd.DataFrame(S["ablation_means"])
    ab = pd.DataFrame(S["ablation"])
    order = [("no_memory", "$-$Memory ($k=1$)"), ("uniform", "$-$Speaker selection (uniform)"), ("round_robin", "$-$Speaker selection (round robin)"),
             ("no_enforce", "$-$Mode enforcement"), ("single_prompt", "Single-prompt baseline")]
    keys = ["responsiveness", "turn_balance", "self_succession", "repetition", "compliant_first", "residual_violation"]
    cols = ["Condition", "Responsiveness", "Turn balance", "Self-succession", "Repetition", "Compliant (1st)", "Residual viol.", "Gap subj.", "Gap FP ratio"]
    if am.empty:
        am = pd.DataFrame(columns=["condition", "reference_for"] + keys)
        ab = pd.DataFrame(columns=["metric", "a", "b", "mean_a", "mean_b", "p"])
    full = am[(am.condition == "full") & (am.reference_for == "no_memory")]
    rows = []
    if len(full):
        gs = ab[(ab.metric == "gapFPminusTP_subjectivity") & (ab.b == "no_memory")]
        gf = ab[(ab.metric == "gapFPminusTP_fp_ratio") & (ab.b == "no_memory")]
        rows.append(["Full framework"] + [fmt(full[k].iloc[0], 3) for k in keys] +
                    [fmt(gs.mean_a.iloc[0], 3) if len(gs) else "--", fmt(gf.mean_a.iloc[0], 3) if len(gf) else "--"])
    for c, lab in order:
        a = am[(am.condition == c) & (am.reference_for == c)]
        if a.empty:
            continue
        gs = ab[(ab.metric == "gapFPminusTP_subjectivity") & (ab.b == c)]
        gf = ab[(ab.metric == "gapFPminusTP_fp_ratio") & (ab.b == c)]
        star = {}
        for k in keys:
            t = ab[(ab.metric == k) & (ab.b == c)]
            star[k] = "$^{*}$" if len(t) and t.p.iloc[0] is not None and np.isfinite(t.p.iloc[0]) and t.p.iloc[0] < 0.05 / len(keys) else ""
        rows.append([lab] + [fmt(a[k].iloc[0], 3) + star[k] for k in keys] +
                    [fmt(gs.mean_b.iloc[0], 3) if len(gs) else "--", fmt(gf.mean_b.iloc[0], 3) if len(gf) else "--"])
    tables[7] = ("Ablation results (Llama 3.1 8B; 6 excerpts $\\times$ 2 modes $\\times$ 3 seeds per condition), compared with the full framework on the same excerpts, modes and seeds.",
                 cols, rows, "lrrrrrrrr",
                 "$^{*}$Paired Wilcoxon test against the full framework, $p<0.05/6$ (Bonferroni over the six measures). Gap: mean First$-$Third difference of subjectivity and of the first-person ratio within a block.")

    # Table 8: judges
    if "judge_by_condition_mode" in S:
        jb = pd.DataFrame(S["judge_by_condition_mode"])
        jb = jb[jb.condition == "full"]
        pa = pd.DataFrame(S.get("judge_perspective_acc", []))
        sid = pd.DataFrame(S.get("speaker_id", []))
        cols = ["Judge", "Mode", "Coherence", "Char. consistency", "Scene fidelity", "Subjectivity", "Emot. intensity", "Form correct", "Speaker ID (chance)"]
        rows = []
        for _, r in jb.sort_values(["judge", "mode"]).iterrows():
            p = pa[(pa.judge == r.judge) & (pa.condition == "full") & (pa["mode"] == r["mode"])] if len(pa) else pd.DataFrame()
            s = sid[(sid.judge == r.judge) & (sid["mode"] == r["mode"])] if len(sid) else pd.DataFrame()
            rows.append([{"phi4:14b": "Phi-4 14B", "qwen3.5:9b-q4_K_M": "Qwen3.5 9B"}.get(r.judge, r.judge), {"first_person": "First", "third_person": "Third", "neutral": "Neutral"}.get(r["mode"], r["mode"]),
                         fmt(r.coherence, 2), fmt(r.character_consistency, 2), fmt(r.scene_fidelity, 2), fmt(r.subjectivity, 2),
                         fmt(r.emotional_intensity, 2), fmt(p.ok.iloc[0], 2) if len(p) else "--",
                         f"{fmt(s.acc.iloc[0], 2)} ({fmt(s.chance.iloc[0], 2)})" if len(s) else "--"])
        tables[8] = ("LLM-judge evaluation of the main-experiment sessions (1--5 scales). Form correct: share of sessions whose dominant narrative form was classified as the instructed one. Speaker ID: accuracy of identifying the speaker of a masked turn.",
                     cols, rows, "llrrrrrrr", None)

    # Table 9: attempts to preserve character identity in third-person narration
    def vs(a, b, d=2):
        return f"{fmt(a, d)} ({fmt(b, d)})" if a is not None and b is not None else "--"

    def judge_pair(rows_, judge, key_a, key_b):
        r = [x for x in rows_ if x.get("judge") == judge]
        return vs(r[0][key_a], r[0][key_b]) if r else "--"

    rows = []
    axm = S.get("attribution_exemplars") or {}
    if axm:
        amx = pd.DataFrame(S.get("ablation_means", []))
        e = amx[(amx.condition == "exemplars") & (amx.reference_for == "exemplars")]
        f0 = amx[(amx.condition == "full") & (amx.reference_for == "exemplars")]
        sx = pd.DataFrame(S.get("speaker_id_exemplars", []))

        def sid_ex(j):
            a = sx[(sx.judge == j) & (sx.condition == "exemplars") & (sx["mode"] == "third_person")]
            b = sx[(sx.judge == j) & (sx.condition == "full") & (sx["mode"] == "third_person")]
            return vs(a.acc.iloc[0], b.acc.iloc[0]) if len(a) and len(b) else "--"
        n_ex = sum(r["n"] for r in S["design"]["by_condition"] if r["condition"] == "exemplars")
        rows.append(["Voice exemplars (both modes)", str(n_ex),
                     vs(e.compliant_first.iloc[0], f0.compliant_first.iloc[0], 3) if len(e) and len(f0) else "--",
                     vs(e.residual_violation.iloc[0], f0.residual_violation.iloc[0], 3) if len(e) and len(f0) else "--",
                     vs(axm["exemplars"]["content"]["cross"], axm["full_matched"]["content"]["cross"]),
                     sid_ex("phi4:14b"), sid_ex("qwen3.5:9b-q4_K_M"), "--", "--"])
    for cond, V, lab in (("voice", "third_person_voice", "Voice-preserving instruction"), ("narrated", "third_person_narrated", "Speak-then-narrate")):
        v = S.get(cond)
        if not v:
            continue
        mm = {r["mode"]: r for r in v["means"]}
        gc = {r["metric"]: r for r in v["gap_change"]}
        ap_ = v["attribution_paired"]["content"]
        sid = v.get("speaker_id", [])
        rows.append([lab, str(v["n_sessions"]),
                     vs(mm[V]["compliant_first"], mm["third_person"]["compliant_first"], 3),
                     vs(mm[V]["residual_violation"], mm["third_person"]["residual_violation"], 3),
                     vs(ap_["cross_voice"], ap_["cross_std"]),
                     judge_pair(sid, "phi4:14b", "acc_voice", "acc_standard"), judge_pair(sid, "qwen3.5:9b-q4_K_M", "acc_voice", "acc_standard"),
                     vs(gc["word_len"]["gap_voice"], gc["word_len"]["gap_standard"]), vs(gc["intensity"]["gap_voice"], gc["intensity"]["gap_standard"], 3)])
    if rows:
        cols = ["Variant", "Sessions", "Compliant (1st)", "Residual viol.", "Cross-mode attribution", "Speaker ID Phi-4", "Speaker ID Qwen3.5",
                "$|$FP$-$TP$|$ word length", "$|$FP$-$TP$|$ intensity"]
        tables[9] = ("Attempts to preserve character identity in third-person narration. Each cell gives the value of the variant and, in parentheses, "
                     "the value of the standard framework on the same excerpts, models and seeds.",
                     cols, rows, "lrrrrrrrr",
                     "Compliance and residual violations refer to third-person turns (exemplars: both modes). Cross-mode attribution: balanced accuracy of a "
                     "content-word classifier trained on one mode and tested on the other (first person vs. the variant's third-person mode). Speaker ID: "
                     "accuracy of the LLM judges on masked third-person turns (seeds 0--2). $|$FP$-$TP$|$: mean absolute first-minus-third difference per "
                     "excerpt, model and seed (content-word length in characters; emotional intensity).")

    for n, (title, cols, rows, align, note) in tables.items():
        tex = latex_table(cols, rows, align)
        if note:
            tex += f"\\\\[2pt]\n{{\\footnotesize {note}}}\n"
        open(os.path.join(tex_dir, f"table{n}.tex"), "w", encoding="utf-8").write(tex)
        docx_table(os.path.join(doc_dir, f"Table{n}.docx"), f"Table {n}. {strip_tex(title)}", cols, rows, note)
        print("table", n, len(rows), "rows")


if __name__ == "__main__":
    main()
